"""Gera o site estático do catálogo Alquimikamente a partir das planilhas e fotos."""
import json, os, re, shutil, unicodedata, difflib
import openpyxl
from PIL import Image, ImageChops

BASE = r"C:/Users/krazm/OneDrive/Área de Trabalho/Catálogo Alquimikamente"
SITE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
for d in ("assets", "img/produtos", "img/fotos"):
    os.makedirs(f"{SITE}/{d}", exist_ok=True)


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().upper()
    return " ".join(re.sub(r"[^A-Z0-9 ]", " ", s).split())


def slug(s):
    return norm(s).lower().replace(" ", "-")


SIGLAS = {"YSL", "EBK", "MFK", "CK", "LV", "D&G", "II"}


def marca(m):
    out = []
    for w in str(m).split():
        if w in SIGLAS:
            out.append(w)
        elif w[:2] in ("D'", "L'"):
            out.append(w[:2] + w[2:].capitalize())
        else:
            out.append(w.capitalize())
    return " ".join(out)


# ---------- perfumes ----------
ws = openpyxl.load_workbook(BASE + "/Fragrâncias Alquimikamente - FINAL.xlsx")["Fragrâncias"]
rows = [r[:5] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
PASTAS = {"FEMININO": "FEMININO", "COMPARTILHÁVEL": "COMPARTILHÁVEL", "MASCULINO": "MASCULINOS"}
GEN = {"FEMININO": "Feminino", "COMPARTILHÁVEL": "Compartilhável", "MASCULINO": "Masculino"}
perfumes, faltando, usados = [], [], set()
for gen, pasta in PASTAS.items():
    fotos = {}
    for f in os.listdir(f"{BASE}/fotos/{pasta}"):
        k = norm(os.path.splitext(f)[0])
        fotos[k[:-2] if k[-2:] in (" F", " M", " C") else k] = f
    for n, g, c, m, fa in sorted(r for r in rows if r[1] == gen):
        k = norm(n)
        f = fotos.get(k)
        if not f:
            alt = difflib.get_close_matches(k, list(fotos), 1, 0.85)
            f = fotos[alt[0]] if alt else next((v for kk, v in fotos.items() if kk.startswith(k + " ") or k.startswith(kk + " ")), None)
        sid = slug(n)
        if sid in usados:
            sid += "-" + slug(GEN[gen])
        usados.add(sid)
        img = ""
        if f:
            im = Image.open(f"{BASE}/fotos/{pasta}/{f}").convert("RGBA")
            flat = Image.alpha_composite(Image.new("RGBA", im.size, (255, 255, 255, 255)), im).convert("RGB")
            diff = ImageChops.difference(flat, Image.new("RGB", flat.size, (255, 255, 255))).convert("L").point(lambda p: 255 if p > 18 else 0)
            flat = flat.crop(diff.getbbox() or (0, 0, *flat.size))
            flat.thumbnail((480, 480))
            img = f"img/produtos/{sid}.webp"
            flat.save(f"{SITE}/{img}", quality=82)
        else:
            faltando.append(n)
        perfumes.append(dict(id=sid, nome=n, genero=GEN[gen], linha=str(c).strip().title(), marca=marca(m), familia=fa or "", img=img))

# ---------- preços (coluna "Preço praticado") ----------
wp = openpyxl.load_workbook(BASE + "/Precificação Alquimikamente.xlsx", data_only=True)["4. Tabela de Preços"]
precos = {}
for r in wp.iter_rows(min_row=5, values_only=True):
    preco = r[7] if r[7] not in (None, "") else r[8]
    if r[0] and r[1] and r[2] and isinstance(preco, (int, float)):
        precos.setdefault(r[0], {}).setdefault(r[1], {})[str(int(r[2]))] = preco

# ---------- fotos de ambiente e logo ----------
for k, src in {"capa": "08_44_21", "sombra": "08_45_09", "categorias": "08_46_27", "spray": "08_47_14"}.items():
    f = [x for x in os.listdir(BASE + "/fotos catálogo") if src in x][0]
    im = Image.open(f"{BASE}/fotos catálogo/{f}").convert("RGB")
    im.thumbnail((1400, 1400))
    im.save(f"{SITE}/img/fotos/{k}.webp", quality=80)
logo = Image.open(BASE + "/Logo Alquimikamente/Logo Alquimikamente.png")
logo.thumbnail((700, 700))
logo.save(f"{SITE}/img/logo.png", optimize=True)
icone = Image.open(BASE + "/Logo Alquimikamente/Logo Alquimikamente - ícone.png")
icone.thumbnail((256, 256))
icone.save(f"{SITE}/img/icone.png", optimize=True)

dados = dict(perfumes=perfumes, precos=precos, atualizado="06/10/2026")
with open(f"{SITE}/assets/dados.js", "w", encoding="utf-8") as fh:
    fh.write("// Gerado automaticamente a partir das planilhas. Não editar à mão.\nwindow.ALQ = ")
    json.dump(dados, fh, ensure_ascii=False, separators=(",", ":"))
    fh.write(";\n")
print("perfumes", len(perfumes), "faltando", faltando)
print(json.dumps(precos, ensure_ascii=False))
